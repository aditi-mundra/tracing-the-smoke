"""Generate formal Microsoft PowerPoint (.pptx) presentation for Tracing the Smoke capstone defense.
Utilizes python-pptx with elegant slide layout, metrics, and structured tables.
"""

import sys
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = Path("E:/AUCML_PROJECT")
EVAL_PATH = PROJECT_ROOT / "audit" / "final" / "FINAL_MODEL_EVALUATION_SUMMARY.json"
DATA_QUAL_PATH = PROJECT_ROOT / "audit" / "final" / "FINAL_DATA_QUALITY_SUMMARY.json"
OUT_PPTX = PROJECT_ROOT / "presentation" / "Tracing_the_Smoke_Presentation.pptx"

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Theme Colors
    c_navy = RGBColor(10, 25, 47)
    c_teal = RGBColor(6, 182, 212)
    c_slate = RGBColor(100, 116, 139)
    c_dark = RGBColor(15, 23, 42)
    c_white = RGBColor(255, 255, 255)
    c_card_bg = RGBColor(241, 245, 249)

    blank_layout = prs.slide_layouts[6]

    # Helper: add header to slide
    def add_header(slide, title_text, category_text="TRACING THE SMOKE — AUCML CAPSTONE"):
        # Category label
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_cat = tb_cat.text_frame.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_teal

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = c_navy

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_navy
    bg1.line.color.rgb = c_navy

    tbox1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(4.0))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "TRACING THE SMOKE"
    p0.font.name = "Arial"
    p0.font.size = Pt(38)
    p0.font.bold = True
    p0.font.color.rgb = c_white

    p1 = tf1.add_paragraph()
    p1.text = "Machine Learning Attribution of Punjab–Haryana Stubble Burning to Delhi's PM2.5 Pollution Episodes"
    p1.font.name = "Arial"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = c_teal
    p1.space_before = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Applications & Use Cases of Machine Learning (AUCML) — ESE Capstone Final Defense\nAcademic Year: 2026 | Verified End-to-End System Evaluation"
    p2.font.name = "Arial"
    p2.font.size = Pt(13)
    p2.font.italic = True
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(24)

    # Slide 2: Real-World Motivation & Research Questions
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Real-World Environmental Motivation & Core Questions")

    box_left = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box_left.fill.solid()
    box_left.fill.fore_color.rgb = c_card_bg
    box_left.line.color.rgb = RGBColor(226, 232, 240)
    tf_l = box_left.text_frame
    tf_l.word_wrap = True
    pl_0 = tf_l.paragraphs[0]
    pl_0.text = "The Air Quality Crisis in Delhi NCT"
    pl_0.font.bold = True
    pl_0.font.size = Pt(15)
    pl_0.font.color.rgb = c_navy

    bullets_l = [
        "Every autumn (Oct–Nov), 33M residents of Delhi NCT suffer catastrophic air pollution breaches (>250 µg/m³ PM2.5).",
        "Public debate is polarized between regional agricultural stubble burning and local urban vehicular/industrial emissions.",
        "Atmospheric advection is non-linear: NW transport winds carry smoke down the Indo-Gangetic Plain, while low nocturnal ventilation traps pollutants near the surface.",
        "Goal: Build an objective, leak-free ML attribution and forecasting engine with verifiable scientific integrity."
    ]
    for b in bullets_l:
        p = tf_l.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = c_dark
        p.space_before = Pt(8)

    box_right = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box_right.fill.solid()
    box_right.fill.fore_color.rgb = c_card_bg
    box_right.line.color.rgb = RGBColor(226, 232, 240)
    tf_r = box_right.text_frame
    tf_r.word_wrap = True
    pr_0 = tf_r.paragraphs[0]
    pr_0.text = "Core Research Questions (RQs)"
    pr_0.font.bold = True
    pr_0.font.size = Pt(15)
    pr_0.font.color.rgb = c_navy

    rqs = [
        ("RQ1 (Prediction Accuracy)", "Can ML reliably predict daily Delhi PM2.5 without lookahead data leakage?"),
        ("RQ2 (Stubble Attribution)", "Does satellite fire data improve accuracy over meteorology-only baselines?"),
        ("RQ3 (Transport Heuristics)", "Do wind-aligned, distance-weighted fire features outperform raw counts?"),
        ("RQ4 (Severe Episodes)", "Can the model accurately detect hazardous episodes (PM2.5 >= 250 µg/m³)?"),
        ("RQ5 (Explainability)", "What physical factors dominate baseline vs acute episodic spikes?")
    ]
    for tag, desc in rqs:
        p = tf_r.add_paragraph()
        p.text = f"• {tag}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = c_dark
        p.space_before = Pt(8)

    # Slide 3: Authentic Multi-Year Datasets (2018–2025)
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Authentic Multi-Year Environmental Datasets (2018–2025)")

    cards = [
        ("CPCB Delhi PM2.5 Network", [
            "40 verified CAAQM stations across Delhi",
            "2,922 continuous calendar days (2018–2025)",
            "2,709 valid daily citywide observations",
            "91,694 total station-day records",
            "157 verified severe episode days (>=250 µg/m³)"
        ]),
        ("NASA FIRMS Satellite Fires", [
            "VIIRS S-NPP (375m resolution)",
            "605,537 authentic detections in Punjab & Haryana",
            "565,728 confirmed cropland detections via ESA WorldCover 2021 (10m)",
            "Indian Standard Time (+5:30) corrected",
            "Accurate FRP (Fire Radiative Power) tracking"
        ]),
        ("ECMWF ERA5 Reanalysis", [
            "2,922 days (100% complete core weather)",
            "Hourly 10m wind speed, direction, u/v vectors",
            "2m Temperature, Dew Point, Relative Humidity",
            "Surface Pressure & Atmospheric Stability",
            "Boundary Layer Height (BLH) & Ventilation Index"
        ])
    ]

    for idx, (title, items) in enumerate(cards):
        left_pos = Inches(0.8 + idx * 3.95)
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.6), Inches(3.75), Inches(5.2))
        box.fill.solid()
        box.fill.fore_color.rgb = c_card_bg
        box.line.color.rgb = RGBColor(203, 213, 225)
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = c_navy

        for item in items:
            p = tf.add_paragraph()
            p.text = "✓ " + item
            p.font.size = Pt(11)
            p.font.color.rgb = c_dark
            p.space_before = Pt(10)

    # Slide 4: Wind-Aware Geospatial Feature Engineering
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Wind-Aware Atmospheric Feature Engineering")

    formulas = [
        ("1. Great-Circle Haversine Distance", "Spherical distance d_i from each fire hotspot to Delhi centroid (R = 6371 km). Accounts for Earth's curvature across the 100–450 km regional transport corridor."),
        ("2. Forward Spherical Bearing", "Exact compass azimuth theta_bearing from fire origin to Delhi using forward spherical trigonometry (typically 135°–160° from Punjab to Delhi)."),
        ("3. Meteorological Wind Conversion", "Meteorological wind alpha denotes direction FROM which wind blows. Smoke advection direction theta_transport = (alpha + 180°) mod 360."),
        ("4. Directional Alignment Factor", "A_i = max(0, cos(theta_bearing - theta_transport)). Zeroes out smoke blowing toward the Himalayas or Pakistan; maximizes downwind smoke aimed at Delhi."),
        ("5. Rational Distance Decay", "w(d_i) = 1 / (1 + (d_i / 150 km)^1.5). Models lateral plume dispersion and dry deposition over multi-hundred kilometer transit."),
        ("6. Multi-Day Advection Lags", "Captures 18–36 hour atmospheric transit delay (fire_upwind_weighted_frp_lag1, lag2, rolling 3d) reflecting realistic 2.5–4.5 m/s transport speeds.")
    ]

    for idx, (f_title, f_desc) in enumerate(formulas):
        col = idx % 2
        row = idx // 2
        l = Inches(0.8 + col * 5.95)
        t = Inches(1.6 + row * 1.75)
        b = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.75), Inches(1.55))
        b.fill.solid()
        b.fill.fore_color.rgb = c_card_bg
        b.line.color.rgb = RGBColor(226, 232, 240)
        tf = b.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = f_title
        pt.font.bold = True
        pt.font.size = Pt(12)
        pt.font.color.rgb = c_navy
        pd = tf.add_paragraph()
        pd.text = f_desc
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = c_dark
        pd.space_before = Pt(4)

    # Slide 5: Chronological Validation & Model Benchmarks Table
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Chronological Model Benchmarks (Held-Out 2024 Test Set, 366 Days)")

    table_shape = s5.shapes.add_table(13, 7, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
    t = table_shape.table

    headers = ["Experiment", "Model", "Features", "MAE (µg/m³)", "RMSE (µg/m³)", "R² Score", "Severe F1"]
    for i, h in enumerate(headers):
        cell = t.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = c_navy
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = c_white
        p.alignment = PP_ALIGN.CENTER

    rows_data = [
        ("Baseline", "Persistence Baseline", "22", "22.24", "37.77", "0.7774", "0.6190"),
        ("Exp A Baseline", "Ridge Regression", "22", "21.76", "35.56", "0.8027", "0.5641"),
        ("Exp A Baseline", "Random Forest", "22", "19.07", "33.66", "0.8233", "0.6111"),
        ("Exp A Baseline", "XGBoost", "22", "18.92", "34.04", "0.8192", "0.5789"),
        ("Exp A Baseline", "LightGBM", "22", "18.48", "33.20", "0.8281", "0.5789"),
        ("Exp B Basic Fires", "Ridge Regression", "27", "21.90", "35.59", "0.8024", "0.5789"),
        ("Exp B Basic Fires", "Random Forest", "27", "18.96", "33.41", "0.8258", "0.5714"),
        ("Exp B Basic Fires", "XGBoost", "27", "18.20", "32.47", "0.8355", "0.6316"),
        ("Exp B Basic Fires", "LightGBM", "27", "18.07", "33.02", "0.8299", "0.6500"),
        ("Exp C Engineered", "Ridge Regression", "38", "22.11", "36.11", "0.7966", "0.5789"),
        ("Exp C Engineered", "Random Forest", "38", "18.70", "32.61", "0.8341", "0.6111"),
        ("Exp C Engineered", "XGBoost (BEST)", "38", "17.86", "31.29", "0.8472", "0.6667")
    ]

    for row_idx, r in enumerate(rows_data, start=1):
        for col_idx, val in enumerate(r):
            cell = t.cell(row_idx, col_idx)
            cell.text = val
            if row_idx == 12:  # Highlight best model
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(224, 242, 254)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            if row_idx == 12:
                p.font.bold = True
                p.font.color.rgb = RGBColor(3, 105, 161)
            p.alignment = PP_ALIGN.CENTER

    # Slide 6: Severe Episode Analysis & Mode B Next-Day Forecasting
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Severe Episodes & Genuine Next-Day Forecasting (Mode B)")

    box_s6_l = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box_s6_l.fill.solid()
    box_s6_l.fill.fore_color.rgb = c_card_bg
    box_s6_l.line.color.rgb = RGBColor(226, 232, 240)
    tfl = box_s6_l.text_frame
    tfl.word_wrap = True
    pl0 = tfl.paragraphs[0]
    pl0.text = "Severe Pollution Detection (PM2.5 >= 250 µg/m³)"
    pl0.font.bold = True
    pl0.font.size = Pt(14)
    pl0.font.color.rgb = c_navy

    b_sev = [
        "2024 held-out test set: 21 verified severe pollution episode days.",
        "Exp A Baseline XGBoost: Detected 11 days (F1 = 0.5789, Recall = 52.4%).",
        "Exp C Engineered XGBoost: Detected 13 days (F1 = 0.6667, Precision = 72.2%, Recall = 61.9%).",
        "Significance: Directional fire transport features reduce false alarms and increase true positive detections of extreme health emergencies.",
        "Zero Fabricated Responses: Evaluated on actual out-of-time test dates without synthetic data."
    ]
    for b in b_sev:
        p = tfl.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = c_dark
        p.space_before = Pt(8)

    box_s6_r = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box_s6_r.fill.solid()
    box_s6_r.fill.fore_color.rgb = c_card_bg
    box_s6_r.line.color.rgb = RGBColor(226, 232, 240)
    tfr = box_s6_r.text_frame
    tfr.word_wrap = True
    pr0 = tfr.paragraphs[0]
    pr0.text = "Mode B: Leakage-Safe Next-Day Forecasting"
    pr0.font.bold = True
    pr0.font.size = Pt(14)
    pr0.font.color.rgb = c_navy

    b_fc = [
        "Predicting day t+1 target PM2.5 strictly using features available at day t.",
        "Persistence Baseline: MAE = 22.20 µg/m³, RMSE = 37.77 µg/m³, R² = 0.7776.",
        "XGBoost Next-Day Forecaster: MAE = 24.31 µg/m³, RMSE = 41.40 µg/m³, R² = 0.7328, Severe Precision = 62.5%.",
        "Scientific Finding: Without future meteorology and local emission forecasts, pure next-day prediction has lower R² (0.7328) than retrospective day-t diagnostic attribution (0.8472).",
        "Inversion & Stagnation Sensitivity: Persistence provides an excellent operational benchmark for day-ahead forecasting."
    ]
    for b in b_fc:
        p = tfr.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = c_dark
        p.space_before = Pt(8)

    # Slide 7: TreeSHAP Explainability & Physical Attribution
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "TreeSHAP Explainability & Physical Attribution")

    shap_table = s7.shapes.add_table(9, 3, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.1))
    st = shap_table.table
    st_hdrs = ["Feature", "Mean |SHAP| (µg/m³)", "Physical Mechanism"]
    for i, h in enumerate(st_hdrs):
        c = st.cell(0, i)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = c_navy
        p = c.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = c_white

    shap_data = [
        ("pm25_lag1", "39.25", "Prior-day pollution persistence (airshed memory)"),
        ("wind_speed_mean", "4.88", "Horizontal ventilation & clearance rate"),
        ("doy_cos", "4.72", "Winter seasonal temperature transition"),
        ("pm25_rolling_3d", "3.77", "Multi-day cumulative particulate build-up"),
        ("temp_min", "3.77", "Nocturnal surface radiation inversion"),
        ("fire_upwind_weighted_frp_lag2", "1.94", "2-day upwind fire transport arrival"),
        ("fire_upwind_weighted_frp_lag1", "1.89", "1-day upwind fire transport arrival"),
        ("fire_count_total (raw)", "0.07", "Raw unweighted fire count (minimal utility)")
    ]
    for row_idx, (f, v, m) in enumerate(shap_data, start=1):
        for col_idx, val in enumerate([f, v, m]):
            c = st.cell(row_idx, col_idx)
            c.text = val
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(9.5)
            if "fire_upwind" in f:
                c.fill.solid()
                c.fill.fore_color.rgb = RGBColor(240, 253, 250)

    # Right text box on SHAP interpretation
    b_shap_r = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.6), Inches(5.3), Inches(5.1))
    b_shap_r.fill.solid()
    b_shap_r.fill.fore_color.rgb = c_card_bg
    b_shap_r.line.color.rgb = RGBColor(226, 232, 240)
    tfs = b_shap_r.text_frame
    tfs.word_wrap = True
    pts = tfs.paragraphs[0]
    pts.text = "Key Scientific Attribution Insights"
    pts.font.bold = True
    pts.font.size = Pt(14)
    pts.font.color.rgb = c_navy

    insights = [
        "Airshed Baseline: Prior-day pollution (39.25 µg/m³) and wind speed (4.88 µg/m³) dictate the general pollution capacity of the NCT.",
        "The Transport Proof: Engineered upwind fire features (lag2: 1.94, lag1: 1.89) have 27x higher attribution than raw counts (0.07).",
        "Acute Peak Driver: During severe November pollution peaks, positive SHAP contributions from upwind fire transport surge to >25% of the total anomaly.",
        "Non-Linear Stagnation: TreeSHAP confirms low wind speed (<2 m/s) acts as a non-linear multiplier on fire smoke impact."
    ]
    for ig in insights:
        p = tfs.add_paragraph()
        p.text = "• " + ig
        p.font.size = Pt(11)
        p.font.color.rgb = c_dark
        p.space_before = Pt(10)

    # Slide 8: Full-Stack Web Application (8 React Pages + FastAPI)
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Production-Grade Full-Stack Web Application")

    app_pages = [
        ("1. Environmental Overview", "Live KPIs, severe pollution alert banner, observed vs predicted time series"),
        ("2. Fire & Transport Map", "Interactive Leaflet map: 40 Delhi stations, active fires, and 150/300 km buffer rings"),
        ("3. Pollution Analysis", "Multi-year historical distributions, seasonal boxplots, station data completeness"),
        ("4. ML Model Laboratory", "Live model benchmarking, actual vs predicted scatter, experiment ablation"),
        ("5. Attribution & SHAP", "Global feature importance bar charts, grouped donut charts, local daily waterfall"),
        ("6. Severe Episode Explorer", "Dual-mode forecasting interface: Mode A (Diagnostic) vs Mode B (Day-Ahead)"),
        ("7. Data & Methodology", "Full dataset dictionary, quality tier table, and mathematical formulations"),
        ("8. Research Insights & SDGs", "Answers to RQ1–RQ5, policy implications, and SDG 3, 11, 13 tracking")
    ]

    for idx, (p_title, p_desc) in enumerate(app_pages):
        col = idx % 2
        row = idx // 2
        l = Inches(0.8 + col * 5.95)
        t = Inches(1.6 + row * 1.3)
        b = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.75), Inches(1.15))
        b.fill.solid()
        b.fill.fore_color.rgb = c_card_bg
        b.line.color.rgb = RGBColor(226, 232, 240)
        tf = b.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = p_title
        pt.font.bold = True
        pt.font.size = Pt(12)
        pt.font.color.rgb = c_navy
        pd = tf.add_paragraph()
        pd.text = p_desc
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = c_dark
        pd.space_before = Pt(2)

    # Slide 9: Software Quality, Automated Testing & Verification
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Software Quality, Automated Testing & Rigorous Verification")

    box_s9_l = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box_s9_l.fill.solid()
    box_s9_l.fill.fore_color.rgb = c_card_bg
    box_s9_l.line.color.rgb = RGBColor(226, 232, 240)
    tf9l = box_s9_l.text_frame
    tf9l.word_wrap = True
    p9l0 = tf9l.paragraphs[0]
    p9l0.text = "Automated Test Suite (100% Pass Rate)"
    p9l0.font.bold = True
    p9l0.font.size = Pt(14)
    p9l0.font.color.rgb = c_navy

    tests_list = [
        "Pytest Suite: 39 automated tests covering end-to-end functionality.",
        "Zero Temporal Leakage: Enforced chronological ordering of train (2018–2022), val (2023), and test (2024).",
        "Geospatial Validation: Great-circle Haversine, forward bearing, wind direction conversions verified.",
        "Model Artifact Verification: Reloading serialized models and running actual live inference.",
        "API Integration: All FastAPI REST endpoints verified with TestClient.",
        "Frontend Build: Built and bundled with TypeScript & Vite (0 errors)."
    ]
    for item in tests_list:
        p = tf9l.add_paragraph()
        p.text = "✓ " + item
        p.font.size = Pt(11)
        p.font.color.rgb = c_dark
        p.space_before = Pt(8)

    box_s9_r = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box_s9_r.fill.solid()
    box_s9_r.fill.fore_color.rgb = c_card_bg
    box_s9_r.line.color.rgb = RGBColor(226, 232, 240)
    tf9r = box_s9_r.text_frame
    tf9r.word_wrap = True
    p9r0 = tf9r.paragraphs[0]
    p9r0.text = "Scientific & Software Verification Deliverables"
    p9r0.font.bold = True
    p9r0.font.size = Pt(14)
    p9r0.font.color.rgb = c_navy

    deliv_list = [
        "Word Academic Report: reports/Tracing_the_Smoke_Academic_Report.docx",
        "Markdown Academic Report: reports/academic_report.md",
        "PowerPoint Defense Presentation: presentation/Tracing_the_Smoke_Presentation.pptx",
        "Viva Defense Guide: presentation/viva_preparation.md",
        "Execution Checkpoint: audit/final/EXECUTION_CHECKPOINT.md",
        "Final Verification Package: audit/final/TRACING_THE_SMOKE_FINAL_VERIFICATION.zip"
    ]
    for item in deliv_list:
        p = tf9r.add_paragraph()
        p.text = "📦 " + item
        p.font.size = Pt(10.5)
        p.font.color.rgb = c_dark
        p.space_before = Pt(8)

    # Slide 10: Conclusion & Defense Summary
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = c_navy
    bg10.line.color.rgb = c_navy

    tbox10 = s10.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.9), Inches(4.8))
    tf10 = tbox10.text_frame
    tf10.word_wrap = True

    p10_0 = tf10.paragraphs[0]
    p10_0.text = "CONCLUSION & PROJECT DEFENSE SUMMARY"
    p10_0.font.name = "Arial"
    p10_0.font.size = Pt(28)
    p10_0.font.bold = True
    p10_0.font.color.rgb = c_white

    takeaways = [
        "Methodological Rigor: Integrating multi-year CPCB ground monitoring, NASA FIRMS active fires, and ECMWF ERA5 reanalysis without lookahead data leakage.",
        "Geospatial Physics: Demonstrating that wind-aligned, distance-weighted fire features provide measurable predictive value (R² = 0.8472 vs 0.8192 without fires).",
        "Operational Value: Accurately detecting 61.9% of severe pollution days with 72.2% precision (F1 = 0.6667) on unseen 2024 test data.",
        "Actionable Governance: Providing local and state authorities with transparent, interpretable TreeSHAP evidence to guide GRAP implementations.",
        "Production Readiness: Fully functional 8-page React dashboard and FastAPI backend backed by a 100% passing automated test suite."
    ]
    for t_item in takeaways:
        p = tf10.add_paragraph()
        p.text = "• " + t_item
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.space_before = Pt(12)

    prs.save(OUT_PPTX)
    print(f"Presentation successfully created at: {OUT_PPTX}")

if __name__ == "__main__":
    create_presentation()
